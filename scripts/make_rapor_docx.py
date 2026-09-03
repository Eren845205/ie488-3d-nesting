"""Markdown rapor -> .docx donusturucu (gorsel destekli).

Kullanim:
  python scripts/make_rapor_docx.py                          # IE488_3D_Nesting_Report.md
  python scripts/make_rapor_docx.py RAPOR_IE488_3D_Nesting.md  # baska dosya

Desteklenen md alt kumesi: #/##/### basliklar, paragraflar, | tablolar,
**bold** satir ici, ![altyazi](yol.png) gorseller, ``` kod bloklari,
- madde isaretleri.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

_ROOT = Path(__file__).resolve().parent.parent


def add_runs(par, text):
    for i, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not piece:
            continue
        run = par.add_run(piece)
        run.bold = (i % 2 == 1)


def main():
    src_name = sys.argv[1] if len(sys.argv) > 1 else "IE488_3D_Nesting_Report.md"
    src = _ROOT / src_name
    dst = src.with_suffix(".docx")

    lines = src.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        m_img = re.match(r"^!\[(.*)\]\((.+)\)\s*$", line.strip())
        if m_img:
            caption, rel_path = m_img.group(1), m_img.group(2)
            img_path = _ROOT / rel_path
            if img_path.exists():
                pic_par = doc.add_paragraph()
                pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic_par.add_run().add_picture(str(img_path), width=Cm(15.5))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(caption)
                run.italic = True
                run.font.size = Pt(9)
            else:
                par = doc.add_paragraph()
                par.add_run(f"[Eksik gorsel: {rel_path}]").italic = True
            i += 1
            continue

        if line.strip().startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # kapanis ```
            par = doc.add_paragraph()
            run = par.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i].strip())
                i += 1
            rows = []
            for tl in tbl_lines:
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if all(set(c) <= {"-", ":", " "} for c in cells):
                    continue
                rows.append(cells)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        par = table.cell(r, c).paragraphs[0]
                        add_runs(par, cell)
                        if r == 0:
                            for run in par.runs:
                                run.bold = True
            continue
        elif line.strip().startswith("- "):
            par = doc.add_paragraph(style="List Bullet")
            add_runs(par, line.strip()[2:])
        else:
            par = doc.add_paragraph()
            add_runs(par, line)
        i += 1

    doc.save(dst)
    print(f"Yazildi: {dst}")


if __name__ == "__main__":
    main()
